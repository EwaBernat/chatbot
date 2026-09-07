#!/usr/bin/env python3
"""Wstawia postac Ewy do prezentacji PowerPoint (.pptx).

Dwa rodzaje wstawki:
  --obraz ewa.png   nieruchoma Ewa z przezroczystym tlem (z wytnij_postac.py --png)
  --klip  ewa.mp4   Ewa mowiaca (film MP4 z jej glosem, np. z wstaw_ewe.py na tle koloru slajdu)

Klip mozna dac ten sam na wybrane slajdy albo inny na kazdy: `--klip 2=wstep.mp4 --klip 5=finał.mp4`.
PowerPoint nie odtwarza filmow z przezroczystoscia, dlatego klip do prezentacji buduj
na tle w kolorze slajdu (wstaw_ewe.py --tlo "#2D1B69" --przytnij) — wtedy wtapia sie w slajd.

Zaleznosci: pip install python-pptx  (do klipu takze ffmpeg — na kadr-miniature).

Przyklady:
    python3 ewa_do_prezentacji.py szkolenie.pptx --obraz ewa.png --slajdy 1,8 --pozycja prawa
    python3 ewa_do_prezentacji.py szkolenie.pptx --klip ewa_wstep.mp4 --slajdy 1 --pozycja srodek
    python3 ewa_do_prezentacji.py szkolenie.pptx --klip 1=wstep.mp4 --klip 6=zakonczenie.mp4
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu, Cm
except ImportError:
    raise SystemExit("Potrzebne: pip install python-pptx")

POZYCJE = ("rog", "lewa", "prawa", "srodek")
WYSOKOSC_DOMYSLNA_CM = {"rog": 6.0, "lewa": 11.0, "prawa": 11.0, "srodek": 14.0}


def znajdz_ffmpeg() -> str | None:
    if szukany := shutil.which("ffmpeg"):
        return szukany
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


def wymiary_obrazu(plik: Path) -> tuple[int, int]:
    from PIL import Image                                   # zaleznosc python-pptx
    with Image.open(plik) as im:
        return im.size


def wymiary_klipu(plik: Path) -> tuple[int, int]:
    try:
        import cv2
        cap = cv2.VideoCapture(str(plik))
        w, h = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        cap.release()
        if w and h:
            return w, h
    except ImportError:
        pass
    ffmpeg = znajdz_ffmpeg()
    if ffmpeg:
        import re
        tekst = subprocess.run([ffmpeg, "-hide_banner", "-i", str(plik)],
                               capture_output=True, text=True).stderr
        m = re.search(r"Video:.*?(\d{2,5})x(\d{2,5})", tekst)
        if m:
            return int(m.group(1)), int(m.group(2))
    return 16, 9


def kadr_miniatury(ffmpeg: str | None, klip: Path, katalog: Path) -> Path | None:
    if not ffmpeg:
        return None
    wyjscie = katalog / f"{klip.stem}_kadr.jpg"
    subprocess.run([ffmpeg, "-hide_banner", "-loglevel", "error", "-y", "-ss", "0.5",
                    "-i", str(klip), "-frames:v", "1", str(wyjscie)], capture_output=True)
    return wyjscie if wyjscie.exists() else None


def zakres_slajdow(tekst: str, ile: int) -> list[int]:
    if tekst in ("wszystkie", "all", "*"):
        return list(range(1, ile + 1))
    numery: set[int] = set()
    for czesc in tekst.split(","):
        czesc = czesc.strip()
        if "-" in czesc:
            a, b = czesc.split("-", 1)
            numery.update(range(int(a), int(b) + 1))
        elif czesc:
            numery.add(int(czesc))
    zle = [n for n in numery if n < 1 or n > ile]
    if zle:
        raise SystemExit(f"Prezentacja ma {ile} slajdow, nie ma numerow: {sorted(zle)}")
    return sorted(numery)


def polozenie(prs, pozycja: str, szer: int, wys: int, margines: int) -> tuple[int, int]:
    W, H = prs.slide_width, prs.slide_height
    if pozycja == "rog":
        return W - szer - margines, H - wys - margines
    if pozycja == "lewa":
        return margines, H - wys
    if pozycja == "prawa":
        return W - szer - margines, H - wys
    return (W - szer) // 2, H - wys                         # srodek, na dole slajdu


def main() -> int:
    ap = argparse.ArgumentParser(description="Wstaw postac Ewy do prezentacji PPTX. Skill awatar-ewa-pctp.")
    ap.add_argument("prezentacja", type=Path)
    ap.add_argument("-o", "--output", type=Path, help="domyslnie <nazwa>_z_ewa.pptx")
    ap.add_argument("--obraz", type=Path, help="PNG z przezroczysta Ewa (wytnij_postac.py --png)")
    ap.add_argument("--klip", action="append", default=[],
                    help="MP4 z Ewa; `plik` na wybrane slajdy albo `N=plik` na slajd N (mozna powtarzac)")
    ap.add_argument("--slajdy", default=None,
                    help="np. 1,3-5 albo wszystkie (domyslnie: obraz -> wszystkie, klip -> 1)")
    ap.add_argument("--pozycja", choices=POZYCJE, default="rog")
    ap.add_argument("--wysokosc-cm", dest="wysokosc_cm", type=float,
                    help="wysokosc Ewy na slajdzie (domyslnie wg pozycji)")
    ap.add_argument("--margines-cm", dest="margines_cm", type=float, default=0.6)
    ap.add_argument("--nazwa", default="Ewa", help="nazwa ksztaltu w PowerPoint")
    a = ap.parse_args()

    if not a.prezentacja.exists():
        print(f"Nie ma pliku: {a.prezentacja}", file=sys.stderr)
        return 1
    if not a.obraz and not a.klip:
        ap.error("podaj --obraz albo --klip")
    if a.obraz and not a.obraz.exists():
        print(f"Nie ma pliku: {a.obraz}", file=sys.stderr)
        return 1

    prs = Presentation(str(a.prezentacja))
    ile = len(prs.slides)
    wys = Cm(a.wysokosc_cm or WYSOKOSC_DOMYSLNA_CM[a.pozycja])
    margines = Cm(a.margines_cm)
    wstawiono: list[str] = []

    # --- obraz -----------------------------------------------------------
    if a.obraz:
        w_px, h_px = wymiary_obrazu(a.obraz)
        szer = int(wys * w_px / h_px)
        for n in zakres_slajdow(a.slajdy or "wszystkie", ile):
            slajd = prs.slides[n - 1]
            x, y = polozenie(prs, a.pozycja, szer, wys, margines)
            ksztalt = slajd.shapes.add_picture(str(a.obraz), Emu(x), Emu(y), Emu(szer), Emu(wys))
            ksztalt.name = a.nazwa
            wstawiono.append(f"slajd {n}: obraz")

    # --- klipy -----------------------------------------------------------
    przydzial: dict[int, Path] = {}
    wspolne: list[Path] = []
    for wpis in a.klip:
        if "=" in wpis and wpis.split("=", 1)[0].strip().isdigit():
            n, plik = wpis.split("=", 1)
            przydzial[int(n)] = Path(plik.strip())
        else:
            wspolne.append(Path(wpis))
    if len(wspolne) > 1:
        ap.error("kilka klipow bez numeru slajdu — uzyj N=plik, zeby przypisac je do slajdow")
    if wspolne:
        for n in zakres_slajdow(a.slajdy or "1", ile):
            przydzial.setdefault(n, wspolne[0])
    for n, plik in przydzial.items():
        if not plik.exists():
            print(f"Nie ma pliku: {plik}", file=sys.stderr)
            return 1
        if n < 1 or n > ile:
            print(f"Prezentacja ma {ile} slajdow, nie ma slajdu {n}", file=sys.stderr)
            return 1

    if przydzial:
        ffmpeg = znajdz_ffmpeg()
        with tempfile.TemporaryDirectory() as tymczasowy:
            for n, plik in sorted(przydzial.items()):
                w_px, h_px = wymiary_klipu(plik)
                szer = int(wys * w_px / h_px)
                x, y = polozenie(prs, a.pozycja, szer, wys, margines)
                miniatura = kadr_miniatury(ffmpeg, plik, Path(tymczasowy))
                typ = "video/mp4" if plik.suffix.lower() in (".mp4", ".m4v") else "video/unknown"
                ksztalt = prs.slides[n - 1].shapes.add_movie(
                    str(plik), Emu(x), Emu(y), Emu(szer), Emu(wys),
                    poster_frame_image=str(miniatura) if miniatura else None, mime_type=typ)
                ksztalt.name = a.nazwa
                wstawiono.append(f"slajd {n}: klip {plik.name}")

    wyjscie = a.output or a.prezentacja.with_name(f"{a.prezentacja.stem}_z_ewa.pptx")
    prs.save(str(wyjscie))
    print(f"Zapisano: {wyjscie}", file=sys.stderr)
    for w in wstawiono:
        print("  " + w, file=sys.stderr)
    if przydzial:
        print("Klip startuje po kliknieciu. Automatyczne odtwarzanie: PowerPoint -> zaznacz Ewe ->\n"
              "Odtwarzanie -> Start: Automatycznie.", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
