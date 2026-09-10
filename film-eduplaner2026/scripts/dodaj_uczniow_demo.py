"""
Dokłada do arkusza „Rejestr" trzech dodatkowych, PRZYKŁADOWYCH uczniów,
kopiując styl pierwszego wiersza danych. Rejestr z jednym wierszem wygląda
na filmie jak pusty szablon — z czterema wygląda jak widok dyrektora.

Uruchomienie: python3 dodaj_uczniow_demo.py ../druki/Baza_Uczniow.xlsx
"""
import sys
from copy import copy

import openpyxl

DEMO = [
    ('UCZ-002', 'Zofia Wiśniewska (dane przykładowe)', '02.11.2015', 'V B', 'szkoła podstawowa · klasy IV–VIII',
     'Nr 88/2026', 'Niepełnosprawność intelektualna w stopniu lekkim', '22.09.2026', '2026/2027', 'mgr Mirosława Jurczyszyn'),
    ('UCZ-003', 'Antoni Zawadzki (dane przykładowe)', '17.03.2019', 'grupa V', 'przedszkole',
     'Nr 12/2026', 'Opóźniony rozwój mowy · ryzyko ASD', '05.10.2026', '2026/2027', 'mgr Mirosława Jurczyszyn'),
    ('UCZ-004', 'Hanna Dąbrowska (dane przykładowe)', '29.06.2013', 'VII A', 'szkoła podstawowa · klasy IV–VIII',
     'Nr 203/2026', 'Zespół Aspergera', '14.10.2026', '2026/2027', 'mgr Mirosława Jurczyszyn'),
]


def main() -> None:
    sciezka = sys.argv[1] if len(sys.argv) > 1 else '../druki/Baza_Uczniow.xlsx'
    wb = openpyxl.load_workbook(sciezka)
    ws = wb['Rejestr']

    # wiersz wzorcowy = pierwszy wiersz danych pod nagłówkiem tabeli
    wzor = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').startswith('UCZ-'):
            wzor = r
            break
    if wzor is None:
        raise SystemExit('nie znalazłem wiersza danych (UCZ-…) w arkuszu Rejestr')

    for i, dane in enumerate(DEMO, start=1):
        docelowy = wzor + i
        ws.row_dimensions[docelowy].height = ws.row_dimensions[wzor].height
        for kol, wartosc in enumerate(dane, start=1):
            zrodlo = ws.cell(wzor, kol)
            cel = ws.cell(docelowy, kol)
            cel.value = wartosc
            cel.font = copy(zrodlo.font)
            cel.fill = copy(zrodlo.fill)
            cel.border = copy(zrodlo.border)
            cel.alignment = copy(zrodlo.alignment)
            cel.number_format = zrodlo.number_format

    wb.save(sciezka)
    print(f'✓ dopisano {len(DEMO)} przykładowych uczniów do {sciezka}')


if __name__ == '__main__':
    main()
